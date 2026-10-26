"""Portable local document operations for the desktop conversation."""
from __future__ import annotations

import hashlib
import io
import json
from pathlib import Path
import re
import shutil
from uuid import uuid4
from xml.sax.saxutils import escape

from sentinel_agent.planner import HybridPlanner

OPERATIONS = {"create_docx", "append_docx", "create_pdf", "append_pdf", "create_xlsx", "edit_xlsx"}


def is_document_goal(goal: str) -> bool:
    return bool(re.search(r"\.(?:docx|xlsx|pdf)\b", goal, re.I))


class DesktopDocuments:
    @staticmethod
    def prepare(action):
        # Some Windows numerical extensions cannot initialize on a worker thread.
        # Load optional format dependencies on the service thread before dispatch.
        if action.name.endswith("xlsx"):
            import openpyxl
        elif action.name.endswith("docx"):
            import docx
        else:
            import pypdf
            import reportlab.platypus

    def __init__(self, root: Path):
        base = root.resolve()
        self.root = base / "artifacts"
        self.root.mkdir(parents=True, exist_ok=True)
        if self.root.resolve() != self.root or self.root.is_symlink():
            raise ValueError("Invalid artifacts directory")

    def path(self, filename: str) -> Path:
        if (Path(filename).name != filename or "/" in filename or "\\" in filename
                or ":" in filename or filename.startswith(".")
                or Path(filename).suffix.lower() not in {".docx", ".xlsx", ".pdf"}):
            raise ValueError("Use a simple .docx, .xlsx or .pdf filename")
        target = self.root / filename
        if target.resolve().parent != self.root or target.is_symlink():
            raise ValueError("Document must stay in the artifacts folder")
        return target

    def plan(self, goal: str):
        if not re.match(r"^\s*(?:create|make|append|add|edit|update|set)\b", goal, re.I):
            raise ValueError("Use create, append, or set with an explicit filename")
        # Reject path-qualified names rather than silently stripping directories.
        if re.search(r"(?:[/\\:]|\.\.)[\w .-]*\.(?:docx|xlsx|pdf)\b", goal, re.I):
            raise ValueError("Use a filename without a directory")
        actions = HybridPlanner._fallback_plan(goal).actions
        if len(actions) != 1 or actions[0].name not in OPERATIONS:
            raise ValueError("Name a .docx, .xlsx or .pdf file")
        action = actions[0]
        self.path(action.arguments["filename"])
        if action.name == "create_xlsx":
            content = re.search(r"\bwith\s+content\s+(.+)$", goal, re.I | re.S)
            if content and content.group(1).lstrip().startswith("{"):
                table = json.loads(content.group(1))
                if not isinstance(table, dict) or set(table) - {"headers", "rows", "sheet_name"}:
                    raise ValueError("Table JSON supports headers, rows and sheet_name")
                action.arguments.update(table)
        return action

    @staticmethod
    def fingerprint(path: Path):
        return hashlib.sha256(path.read_bytes()).hexdigest() if path.exists() else None

    def execute(self, action, fingerprint):
        name, args = action.name, action.arguments
        if name not in OPERATIONS:
            raise ValueError("Unsupported document operation")
        target = self.path(args["filename"])
        if self.fingerprint(target) != fingerprint:
            raise ValueError("File changed after review; submit the task again")
        if name.startswith(("append_", "edit_")) and not target.is_file():
            raise ValueError("Import or create the document first")
        temporary = self.root / f".{uuid4().hex}{target.suffix}"
        try:
            if target.suffix.lower() == ".docx":
                from docx import Document
                document = Document(target) if name == "append_docx" else Document()
                if name == "create_docx":
                    document.add_heading(args.get("title", target.stem), 0)
                for paragraph in args.get("paragraphs", []):
                    document.add_paragraph(str(paragraph))
                document.save(temporary)
                verified = Document(temporary)
                if not verified.paragraphs:
                    raise ValueError("Word verification failed")
            elif target.suffix.lower() == ".pdf":
                from pypdf import PdfReader, PdfWriter
                from reportlab.lib.styles import getSampleStyleSheet
                from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
                styles = getSampleStyleSheet()
                addition = io.BytesIO()
                story = [Paragraph(escape(str(args.get("title", target.stem))), styles["Title"])]
                for paragraph in args.get("paragraphs", []):
                    story.extend([Paragraph(escape(str(paragraph)).replace("\n", "<br/>"), styles["BodyText"]), Spacer(1, 10)])
                SimpleDocTemplate(addition).build(story)
                writer = PdfWriter()
                if name == "append_pdf":
                    writer.append(PdfReader(target))
                writer.append(PdfReader(addition))
                with temporary.open("wb") as stream:
                    writer.write(stream)
                if not PdfReader(temporary).pages:
                    raise ValueError("PDF verification failed")
            else:
                from openpyxl import Workbook, load_workbook
                from openpyxl.styles import Font, PatternFill
                workbook = load_workbook(target) if name == "edit_xlsx" else Workbook()
                sheet_name = args.get("sheet_name", "Sheet1")
                if name == "create_xlsx":
                    sheet = workbook.active
                    sheet.title = sheet_name
                    headers, rows = args.get("headers", ["Item", "Value"]), args.get("rows", [])
                    if not isinstance(headers, list) or not isinstance(rows, list) or len(rows) > 5000:
                        raise ValueError("Invalid or oversized table")
                    sheet.append(headers)
                    for row in rows:
                        if not isinstance(row, list) or len(row) > 200:
                            raise ValueError("Invalid table row")
                        sheet.append(row)
                    for cell in sheet[1]:
                        cell.font = Font(bold=True, color="FFFFFF")
                        cell.fill = PatternFill("solid", fgColor="334155")
                    sheet.freeze_panes = "A2"
                    sheet.auto_filter.ref = sheet.dimensions
                    for column in sheet.columns:
                        sheet.column_dimensions[column[0].column_letter].width = min(55, max(14, max(len(str(c.value or "")) for c in column) + 2))
                else:
                    sheet = workbook[sheet_name]
                    for update in args.get("updates", []):
                        if not re.fullmatch(r"[A-Z]{1,3}[1-9]\d{0,6}", update["cell"]):
                            raise ValueError("Invalid cell")
                        sheet[update["cell"]] = update.get("formula", update.get("value"))
                workbook.save(temporary)
                workbook.close()
                verified = load_workbook(temporary)
                if sheet_name not in verified.sheetnames:
                    raise ValueError("Workbook verification failed")
                verified.close()
            # Recheck after generation, then back up before replacing an existing file.
            if self.fingerprint(target) != fingerprint:
                raise ValueError("File changed during generation")
            backup = None
            if target.exists():
                backup_dir = self.root / "backups"
                backup_dir.mkdir(exist_ok=True)
                if backup_dir.resolve().parent != self.root:
                    raise ValueError("Invalid backup directory")
                backup = backup_dir / f"{target.stem}-{uuid4().hex}{target.suffix}"
                shutil.copy2(target, backup)
            temporary.replace(target)
            return {"filename": target.name, "path": str(target), "bytes": target.stat().st_size,
                    "backup_path": str(backup) if backup else None, "verified": True}
        finally:
            temporary.unlink(missing_ok=True)
