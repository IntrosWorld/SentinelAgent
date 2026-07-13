from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader, PdfWriter
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


def create_docx(path: Path, title: str, paragraphs: list[str]) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.85)
    section.left_margin = section.right_margin = Inches(1)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_paragraph.add_run(title)
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(23, 37, 84)
    title_paragraph.paragraph_format.space_after = Pt(18)
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(path)


def create_pdf(path: Path, title: str, paragraphs: list[str]) -> None:
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "SentinelTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=22,
        leading=27,
        alignment=TA_CENTER,
        textColor=HexColor("#172554"),
        spaceAfter=20,
    )
    body_style = ParagraphStyle(
        "SentinelBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10.5,
        leading=15,
        textColor=HexColor("#1F2937"),
        spaceAfter=9,
    )
    doc = SimpleDocTemplate(
        str(path), pagesize=LETTER, rightMargin=inch, leftMargin=inch, topMargin=0.8 * inch, bottomMargin=0.8 * inch
    )
    story = [Paragraph(title, title_style)]
    for text in paragraphs:
        story.extend([Paragraph(text, body_style), Spacer(1, 2)])
    doc.build(story)


payload = json.loads(sys.argv[1])
operation = payload["operation"]
output = Path(payload["output_path"])
output.parent.mkdir(parents=True, exist_ok=True)

if operation == "create_docx":
    create_docx(output, payload["title"], payload.get("paragraphs", []))
elif operation == "append_docx":
    document = Document(output)
    for text in payload.get("paragraphs", []):
        document.add_paragraph(text)
    document.save(output)
elif operation == "create_pdf":
    create_pdf(output, payload["title"], payload.get("paragraphs", []))
elif operation == "append_pdf":
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as stream:
        addition = Path(stream.name)
    create_pdf(addition, payload["title"], payload.get("paragraphs", []))
    writer = PdfWriter()
    writer.append(PdfReader(output))
    writer.append(PdfReader(addition))
    with output.open("wb") as stream:
        writer.write(stream)
    addition.unlink(missing_ok=True)
else:
    raise ValueError(f"Unsupported operation: {operation}")

if output.suffix.lower() == ".docx":
    verification = Document(output)
    text = "\n".join(paragraph.text for paragraph in verification.paragraphs)
    metadata = {"output_path": str(output), "paragraphs": len(verification.paragraphs), "characters": len(text)}
else:
    verification = PdfReader(output)
    text = "\n".join(page.extract_text() or "" for page in verification.pages)
    metadata = {"output_path": str(output), "pages": len(verification.pages), "characters": len(text)}
print(json.dumps(metadata))

