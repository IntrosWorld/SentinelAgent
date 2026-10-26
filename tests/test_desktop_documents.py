import asyncio
from pathlib import Path

import pytest

from sentinel_agent.desktop_documents import DesktopDocuments
from sentinel_agent.desktop_service import DesktopService


async def run_document(service, events, goal, approve=True):
    events.clear()
    await service.dispatch("run.start", {"mode": "auto", "goal": goal})
    async def drive():
        while not service.running.done():
            if service.pending and not service.pending[1].done():
                await service.dispatch("run.approve", {
                    "run_id": service.run_id, "token": service.pending[0], "approved": approve})
            await asyncio.sleep(.01)
        await service.running
    await asyncio.wait_for(drive(), 20)
    return next(e["result"] for e in events if e["event"] == "finished")


async def test_real_documents_through_service_create_edit_and_verify(tmp_path):
    from docx import Document
    from openpyxl import load_workbook
    from pypdf import PdfReader
    events = []
    service = DesktopService(tmp_path, events.append)
    for goal in [
        "create report.docx with content First paragraph",
        "append report.docx with content Second paragraph",
        "create report.pdf with content First page",
        "append report.pdf with content Second page",
        'create results.xlsx with content {"headers":["Item","Value"],"rows":[["Sales",42]]}',
        "set B3 to 99 in results.xlsx",
    ]:
        result = await run_document(service, events, goal)
        assert result["status"] == "completed", result
        assert result["artifacts"][0]["verified"]
        assert any(e["event"] == "approval" for e in events)
    artifacts = tmp_path / "artifacts"
    assert "Second paragraph" in "\n".join(p.text for p in Document(artifacts / "report.docx").paragraphs)
    assert len(PdfReader(artifacts / "report.pdf").pages) == 2
    workbook = load_workbook(artifacts / "results.xlsx")
    assert workbook.active["B2"].value == 42
    assert workbook.active["B3"].value == 99
    workbook.close()
    assert len(list((artifacts / "backups").iterdir())) == 3
    assert service.history[-1]["artifacts"][0]["filename"] == "results.xlsx"


async def test_decline_and_stop_never_create_files(tmp_path):
    events = []
    service = DesktopService(tmp_path, events.append)
    result = await run_document(service, events, "create declined.docx with content Test", approve=False)
    assert result["status"] == "cancelled"
    assert not (tmp_path / "artifacts" / "declined.docx").exists()
    await service.dispatch("run.start", {"mode": "documents", "goal": "create stopped.pdf with content Test"})
    while service.pending is None:
        await asyncio.sleep(.01)
    await service.dispatch("run.cancel", {"run_id": service.run_id})
    await service.running
    assert not (tmp_path / "artifacts" / "stopped.pdf").exists()


@pytest.mark.parametrize("goal", ["create ../outside.pdf", "create C:\\outside.docx", "delete report.pdf", "create hello.txt"])
def test_invalid_document_commands(tmp_path, goal):
    with pytest.raises(ValueError):
        DesktopDocuments(tmp_path).plan(goal)


def test_reviewed_file_change_is_rejected(tmp_path):
    documents = DesktopDocuments(tmp_path)
    action = documents.plan("create report.docx with content Test")
    target = documents.path("report.docx")
    target.write_bytes(b"external change")
    with pytest.raises(ValueError, match="changed"):
        documents.execute(action, None)
    assert target.read_bytes() == b"external change"


def test_content_verbs_and_multiple_lines_do_not_change_operation(tmp_path):
    documents = DesktopDocuments(tmp_path)
    action = documents.plan("create report.docx with content Add details\nUpdate the summary")
    assert action.name == "create_docx"
    assert action.arguments["paragraphs"] == ["Add details\nUpdate the summary"]


async def test_missing_edit_is_failed_without_approval(tmp_path):
    events = []
    service = DesktopService(tmp_path, events.append)
    result = await run_document(service, events, "append missing.docx with content Text")
    assert result["status"] == "failed"
    assert not any(e["event"] == "approval" for e in events)
